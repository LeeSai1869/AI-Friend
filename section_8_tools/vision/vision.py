"""Local document, image, and video reading tools."""
from __future__ import annotations

import base64
import binascii
from io import BytesIO
import json
import math
import re
import shutil
import subprocess
import tempfile
import urllib.request
from pathlib import Path
from typing import Any, Dict

from _base.tool import Tool
from _base.paths import missing_parent_directories
from guard.contract import ToolEffect


OLLAMA_URL = "http://localhost:11434"
VISION_MODEL = "qwen3-vl:8b"
VIDEO_EXTENSIONS = {".avi", ".m4v", ".mkv", ".mov", ".mp4", ".mpeg", ".mpg", ".webm"}
VIDEO_EXTENSIONS.update({
    ".3g2", ".3gp", ".asf", ".f4v", ".flv", ".m2ts", ".mts",
    ".ts", ".vob", ".wmv",
})
IMAGE_EXTENSIONS = {
    ".bmp", ".gif", ".heic", ".ico", ".jpeg", ".jpg", ".png",
    ".tif", ".tiff", ".webp",
}
VISION_DOCUMENT_EXTENSIONS = IMAGE_EXTENSIONS | {".docx", ".pdf", ".pptx"}


def _ffmpeg_executable() -> str:
    try:
        import imageio_ffmpeg
    except ImportError:
        executable = shutil.which("ffmpeg")
        if executable:
            return executable
        raise RuntimeError(
            "Video reading requires the bundled imageio-ffmpeg runtime dependency."
        )
    return imageio_ffmpeg.get_ffmpeg_exe()


def _ocr_image_bytes(
    image_bytes: bytes,
    prompt: str | None = None,
    num_predict: int = 4096,
) -> str:
    try:
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Image reading requires the Pillow runtime dependency.") from exc

    try:
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            pass
        with Image.open(BytesIO(image_bytes)) as image:
            replies = []
            frame_count = getattr(image, "n_frames", 1)
            for index in range(frame_count):
                image.seek(index)
                image.load()
                normalized = BytesIO()
                image.convert("RGB").save(normalized, format="PNG")
                image_b64 = base64.b64encode(normalized.getvalue()).decode("ascii")
                reply = _call_vision(
                    prompt or (
                        "Read this image carefully. Return two sections: "
                        "VISIBLE TEXT (transcribe visible text faithfully, preserving "
                        "line breaks and original scripts), and VISUAL DESCRIPTION "
                        "(describe the screen, objects, layout, and relevant context). "
                        "Do not guess text that is not legible."
                    ),
                    image_b64,
                    num_predict=num_predict,
                )
                replies.append(
                    f"----- image frame {index + 1} -----\n{reply}"
                    if frame_count > 1
                    else reply
                )
    except Exception as exc:
        if isinstance(exc, (RuntimeError, OSError)):
            raise
        raise ValueError(f"could not decode image: {exc}") from exc

    return "\n\n".join(replies)


def _read_pdf_bytes(data: bytes, pages_arg: str = "all") -> str:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise RuntimeError("PDF reading requires the PyMuPDF runtime dependency.") from exc

    pages = []
    try:
        document = fitz.open(stream=data, filetype="pdf")
        if document.needs_pass:
            raise ValueError("PDF is password protected; provide an unlocked PDF.")
        if pages_arg == "all":
            page_range = range(len(document))
        else:
            match = re.match(r"^(\d+)-(\d+)$", pages_arg)
            if not match:
                raise ValueError("pages must be 'all' or 'N-M'.")
            start, end = (int(value) for value in match.groups())
            if start < 1 or end < start or start > len(document):
                raise ValueError("page range is outside the document.")
            page_range = range(start - 1, min(end, len(document)))
        for index in page_range:
            page_number = index + 1
            page = document[index]
            text = page.get_text("text").strip()
            if len(text) < 30:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_b64 = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
                ocr_text = _call_vision(
                    "Transcribe all visible text on this PDF page. Preserve "
                    "the original script and layout as closely as possible. "
                    "Do not add commentary or invent unreadable text.",
                    image_b64,
                    num_predict=4096,
                ).strip()
                text = f"[OCR]\n{ocr_text}" if ocr_text else "[No text could be recognized on this page.]"
            pages.append(f"----- page {page_number} -----\n{text}")
    except (OSError, ValueError, RuntimeError):
        raise
    except Exception as exc:
        raise ValueError(f"could not read PDF: {exc}") from exc
    finally:
        if "document" in locals():
            document.close()

    if not pages:
        raise ValueError("PDF contains no pages.")
    return "\n\n".join(pages)


def _read_docx_bytes(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Word document reading requires python-docx.") from exc
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            parts.append(f"[Table {table_index}]\n" + "\n".join(rows))
        for related_part in document.part.related_parts.values():
            if getattr(related_part, "content_type", "").startswith("image/"):
                image_text = _ocr_image_bytes(related_part.blob).strip()
                if image_text:
                    parts.append(f"[Embedded image OCR and description]\n{image_text}")
        return "\n\n".join(parts)
    except (OSError, RuntimeError):
        raise
    except Exception as exc:
        raise ValueError(f"could not read Word document: {exc}") from exc


def _read_pptx_bytes(data: bytes) -> str:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError("PowerPoint reading requires python-pptx.") from exc
    try:
        presentation = Presentation(BytesIO(data))
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    parts.extend(
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                    )
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    image_text = _ocr_image_bytes(shape.image.blob).strip()
                    if image_text:
                        parts.append(f"[Image OCR and description]\n{image_text}")
            slides.append(f"----- slide {slide_number} -----\n" + "\n".join(parts))
        return "\n\n".join(slides)
    except (OSError, RuntimeError):
        raise
    except Exception as exc:
        raise ValueError(f"could not read PowerPoint presentation: {exc}") from exc


def _read_xlsx_bytes(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel workbook reading requires openpyxl.") from exc
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append(" | ".join(values))
            sheets.append(f"----- sheet {sheet.title} -----\n" + "\n".join(rows))
        workbook.close()
        return "\n\n".join(sheets)
    except Exception as exc:
        raise ValueError(f"could not read Excel workbook: {exc}") from exc


def read_document_bytes(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _read_pdf_bytes(data)
    if suffix in IMAGE_EXTENSIONS:
        return _ocr_image_bytes(data)
    if suffix == ".docx":
        return _read_docx_bytes(data)
    if suffix == ".pptx":
        return _read_pptx_bytes(data)
    if suffix == ".xlsx":
        return _read_xlsx_bytes(data)
    if suffix in VIDEO_EXTENSIONS:
        raise ValueError("This is a video file; use the video reader.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(
            f'"{name}" is not a supported document or readable text file.'
        ) from exc
    if "\x00" in text:
        raise ValueError(
            f'"{name}" is not a supported document or readable text file.'
        )
    return text


def read_video_bytes(name: str, data: bytes, fps: float = 1.0) -> str:
    if Path(name).suffix.lower() not in VIDEO_EXTENSIONS:
        raise ValueError(f'"{name}" is not a supported video format.')
    if not math.isfinite(fps) or fps <= 0:
        raise ValueError("fps must be a positive finite number.")
    try:
        executable = _ffmpeg_executable()
    except (RuntimeError, OSError) as exc:
        raise RuntimeError(str(exc)) from exc

    with tempfile.TemporaryDirectory(prefix="ai-friend-video-") as temp_dir:
        source = Path(temp_dir) / Path(name).name
        source.write_bytes(data)
        output_pattern = Path(temp_dir) / "frame-%08d.jpg"
        command = [
            executable, "-nostdin", "-v", "error", "-i", str(source),
            "-vf", f"fps={fps:.6g}", "-vsync", "0",
            str(output_pattern),
        ]
        process = subprocess.run(
            command,
            capture_output=True,
            text=True,
        )
        if process.returncode:
            detail = process.stderr.strip()[-1000:]
            raise ValueError(f"could not decode video: {detail or 'ffmpeg failed'}")

        frames = sorted(Path(temp_dir).glob("frame-*.jpg"))
        if not frames:
            raise ValueError("video contains no decodable frames.")
        parts = []
        for index, frame in enumerate(frames):
            timestamp = index / fps
            description = _ocr_image_bytes(
                frame.read_bytes(),
                "Describe this video frame at timestamp "
                f"{timestamp:.2f} seconds. Transcribe visible on-screen text "
                "faithfully, then describe the visible scene and actions. "
                "Do not infer events outside this frame.",
                num_predict=1024,
            )
            parts.append(f"----- frame at {timestamp:.2f}s -----\n{description}")
        return "\n\n".join(parts)


def _encode_image_b64(path: Path) -> str:
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


def _call_vision(
    prompt: str,
    image_b64: str,
    timeout: int = 180,
    num_predict: int = 1024,
) -> str:
    payload = {
        "model": VISION_MODEL,
        "messages": [
            {
                "role": "user",
                "content": prompt,
                "images": [image_b64],
            }
        ],
        "stream": False,
        "options": {"num_predict": num_predict},
    }
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{OLLAMA_URL}/api/chat",
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        parsed = json.loads(resp.read().decode("utf-8"))
        return parsed.get("message", {}).get("content", "")


class ImageDescribe(Tool):
    name = "image_describe"
    description = (
        "Look at an image and describe it, or answer a question "
        "about it. Uses the local vision model."
    )
    parameters = {
        "path": "path to the image file (relative to workspace or absolute)",
        "question": "optional question about the image",
    }

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        rel = (args.get("path") or "").strip()
        if not rel:
            raise ValueError("path is required")
        target = self._resolve(rel)

        return (
            ToolEffect(
                action="read",
                target=str(target),
                scope="exact",
            ),
            ToolEffect(
                action="service",
                target=f"ollama:{VISION_MODEL}",
                scope="service",
            ),
        )

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        rel = (args.get("path") or "").strip()
        question = (args.get("question") or "").strip()

        if not rel:
            return {"ok": False, "data": None, "error": "path is required"}

        try:
            target = self._resolve(rel)
        except ValueError as e:
            return {"ok": False, "data": None, "error": str(e)}

        if not target.exists():
            return {"ok": False, "data": None, "error": f"not found: {rel}"}

        try:
            image_b64 = _encode_image_b64(target)
        except Exception as e:
            return {"ok": False, "data": None, "error": f"read failed: {e}"}

        prompt = question or "Describe this image in detail. What do you see?"
        try:
            reply = _call_vision(prompt, image_b64)
        except Exception as e:
            return {"ok": False, "data": None, "error": f"vision model error: {e}"}

        return {
            "ok": True,
            "data": {
                "path": str(target),
                "question": question or "(describe)",
                "answer": reply,
            },
            "error": None,
        }


class ImageReadText(Tool):
    name = "image_read_text"
    description = (
        "Read text from an image (OCR). Handles screenshots, scans, "
        "and photographed documents. Works for any script, including "
        "Urdu, Arabic, Chinese, Hindi."
    )
    parameters = {
        "path": "path to the image file",
    }

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        rel = (args.get("path") or "").strip()
        if not rel:
            raise ValueError("path is required")
        target = self._resolve(rel)

        return (
            ToolEffect(
                action="read",
                target=str(target),
                scope="exact",
            ),
            ToolEffect(
                action="service",
                target=f"ollama:{VISION_MODEL}",
                scope="service",
            ),
        )

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        rel = (args.get("path") or "").strip()
        if not rel:
            return {"ok": False, "data": None, "error": "path is required"}

        try:
            target = self._resolve(rel)
        except ValueError as e:
            return {"ok": False, "data": None, "error": str(e)}

        if not target.exists():
            return {"ok": False, "data": None, "error": f"not found: {rel}"}

        try:
            image_b64 = _encode_image_b64(target)
        except Exception as e:
            return {"ok": False, "data": None, "error": f"read failed: {e}"}

        prompt = (
            "Extract all text visible in this image. Return only the "
            "text, preserving line breaks and structure. Do not add "
            "commentary. If the text is in a non-Latin script, keep "
            "it in its original script and also provide a "
            "transliteration on the next line."
        )
        try:
            reply = _call_vision(prompt, image_b64)
        except Exception as e:
            return {"ok": False, "data": None, "error": f"vision model error: {e}"}

        return {
            "ok": True,
            "data": {
                "path": str(target),
                "text": reply,
            },
            "error": None,
        }


class PdfRead(Tool):
    name = "pdf_read"
    description = (
        "Extract embedded PDF text and OCR scanned pages using the local "
        "vision model."
    )
    parameters = {
        "path": "path to the PDF file",
        "pages": "optional page range, e.g. '1-5' or 'all'",
    }

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        rel = (args.get("path") or "").strip()
        if not rel:
            raise ValueError("path is required")
        target = self._resolve(rel)

        return (
            ToolEffect(
                action="read",
                target=str(target),
                scope="exact",
            ),
            ToolEffect(
                action="service",
                target=f"ollama:{VISION_MODEL}",
                scope="service",
            ),
        )

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        rel = (args.get("path") or "").strip()
        if not rel:
            return {"ok": False, "data": None, "error": "path is required"}

        try:
            target = self._resolve(rel)
        except ValueError as e:
            return {"ok": False, "data": None, "error": str(e)}

        if not target.exists():
            return {"ok": False, "data": None, "error": f"not found: {rel}"}

        pages_arg = (args.get("pages") or "all").strip()
        try:
            full_text = _read_pdf_bytes(target.read_bytes(), pages_arg)
        except (OSError, RuntimeError, ValueError) as exc:
            return {"ok": False, "data": None, "error": str(exc)}

        return {
            "ok": True,
            "data": {
                "path": str(target),
                "pages_read": full_text.count("----- page "),
                "text": full_text,
            },
            "error": None,
        }


class DocumentRead(Tool):
    name = "document_read"
    description = (
        "Read text documents, PDF files (including scanned-page OCR), "
        "DOCX, PPTX, and XLSX files, and images or screenshots. "
        "Uploads can be passed as base64 without saving them to the workspace."
    )
    parameters = {
        "path": "workspace-relative path to a document",
        "name": "original filename when data_base64 is provided",
        "data_base64": "base64-encoded uploaded file bytes",
    }

    def _input(self, args: Dict[str, Any]) -> tuple[str, bytes, Path | None]:
        encoded = args.get("data_base64")
        if encoded is not None:
            if not isinstance(encoded, str):
                raise ValueError("data_base64 must be a string.")
            name = str(args.get("name") or "uploaded-document")
            try:
                return name, base64.b64decode(encoded, validate=True), None
            except (binascii.Error, ValueError) as exc:
                raise ValueError("uploaded document is not valid base64.") from exc
        rel = str(args.get("path") or "").strip()
        if not rel:
            raise ValueError("path or data_base64 is required.")
        target = self._resolve(rel)
        if not target.is_file():
            raise ValueError(f"not a file: {rel}")
        return target.name, target.read_bytes(), target

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        effects = []
        if args.get("data_base64") is None:
            rel = str(args.get("path") or "").strip()
            if not rel:
                raise ValueError("path or data_base64 is required.")
            target = self._resolve(rel)
            effects.append(ToolEffect(action="read", target=str(target)))
            name = target.name
        else:
            name = str(args.get("name") or "")
        effects.append(ToolEffect(
            action="service",
            target="document-reader",
            scope="service",
        ))
        if Path(name).suffix.lower() in VISION_DOCUMENT_EXTENSIONS:
            effects.append(ToolEffect(
                action="service",
                target=f"ollama:{VISION_MODEL}",
                scope="service",
            ))
        return tuple(effects)

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        try:
            name, data, path = self._input(args)
            text = read_document_bytes(name, data)
        except (OSError, RuntimeError, ValueError) as exc:
            return {"ok": False, "data": None, "error": str(exc)}
        return {
            "ok": True,
            "data": {
                "name": name,
                "path": str(path) if path else None,
                "text": text,
                "content_bytes": len(text.encode("utf-8")),
            },
            "error": None,
        }


class VideoRead(Tool):
    name = "video_read"
    description = (
        "Analyze a video by sampling every requested interval, describing "
        "each frame, and transcribing visible on-screen text. Uses local "
        "ffmpeg and the local vision model; no frame-count cap is applied."
    )
    parameters = {
        "path": "workspace-relative path to a video",
        "name": "original filename when data_base64 is provided",
        "data_base64": "base64-encoded uploaded video bytes",
        "fps": "optional frames sampled per second; defaults to 1",
    }

    def _input(self, args: Dict[str, Any]) -> tuple[str, bytes, Path | None]:
        encoded = args.get("data_base64")
        if encoded is not None:
            if not isinstance(encoded, str):
                raise ValueError("data_base64 must be a string.")
            name = str(args.get("name") or "uploaded-video")
            try:
                return name, base64.b64decode(encoded, validate=True), None
            except (binascii.Error, ValueError) as exc:
                raise ValueError("uploaded video is not valid base64.") from exc
        rel = str(args.get("path") or "").strip()
        if not rel:
            raise ValueError("path or data_base64 is required.")
        target = self._resolve(rel)
        if not target.is_file():
            raise ValueError(f"not a file: {rel}")
        return target.name, target.read_bytes(), target

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        effects = []
        if args.get("data_base64") is None:
            rel = str(args.get("path") or "").strip()
            if not rel:
                raise ValueError("path or data_base64 is required.")
            effects.append(ToolEffect(action="read", target=str(self._resolve(rel))))
        effects.extend((
            ToolEffect(action="execute", target="ffmpeg", scope="service"),
            ToolEffect(
                action="service",
                target=f"ollama:{VISION_MODEL}",
                scope="service",
            ),
        ))
        return tuple(effects)

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        try:
            name, data, path = self._input(args)
            raw_fps = args.get("fps", 1)
            fps = float(1 if raw_fps in (None, "") else raw_fps)
            text = read_video_bytes(name, data, fps)
        except (OSError, RuntimeError, TypeError, ValueError) as exc:
            return {"ok": False, "data": None, "error": str(exc)}
        return {
            "ok": True,
            "data": {
                "name": name,
                "path": str(path) if path else None,
                "text": text,
                "sampled_fps": fps,
            },
            "error": None,
        }


class VideoExtractFrames(Tool):
    name = "video_extract_frames"
    description = (
        "Extract frames from a video as PNG images. Each frame can "
        "then be passed to image_describe or image_read_text."
    )
    parameters = {
        "path": "path to the video file",
        "fps": "optional frames per second, default 1",
        "max_frames": "optional cap, default 30",
    }

    def effects(self, args: Dict[str, Any]) -> tuple[ToolEffect, ...]:
        rel = (args.get("path") or "").strip()
        if not rel:
            raise ValueError("path is required")

        target = self._resolve(rel)
        out_dir = (
            self.workspace_root
            / ".central_mind"
            / "frames"
        ).resolve()

        effects = [
            ToolEffect(
                action="read",
                target=str(target),
                scope="exact",
            ),
        ]

        # mkdir(parents=True) may create ancestors of frames. Declare
        # every such mutation explicitly; tree scope never reaches upward.
        effects.extend(
            ToolEffect(
                action="create",
                target=str(parent),
                scope="exact",
            )
            for parent in missing_parent_directories(
                self.workspace_root,
                out_dir,
            )
        )

        if not out_dir.exists():
            effects.append(
                ToolEffect(
                    action="create",
                    target=str(out_dir),
                    scope="exact",
                )
            )

        # ffmpeg -y may create new frame files or overwrite existing
        # matching frame files below this directory.
        effects.append(
            ToolEffect(
                action="write",
                target=str(out_dir),
                scope="tree",
            )
        )

        effects.append(
            ToolEffect(
                action="execute",
                target="ffmpeg",
                scope="service",
            )
        )

        return tuple(effects)

    def call(self, args: Dict[str, Any]) -> Dict[str, Any]:
        rel = (args.get("path") or "").strip()
        if not rel:
            return {"ok": False, "data": None, "error": "path is required"}

        try:
            target = self._resolve(rel)
        except ValueError as e:
            return {"ok": False, "data": None, "error": str(e)}

        if not target.exists():
            return {"ok": False, "data": None, "error": f"not found: {rel}"}

        fps = float(args.get("fps") or 1)
        max_frames = int(args.get("max_frames") or 30)

        out_dir = self.workspace_root / ".central_mind" / "frames"
        out_dir.mkdir(parents=True, exist_ok=True)

        pattern = str(out_dir / f"{target.stem}_%04d.png")

        try:
            ffmpeg = _ffmpeg_executable()
        except (RuntimeError, OSError) as e:
            return {"ok": False, "data": None, "error": str(e)}
        cmd = [
            ffmpeg,
            "-y",
            "-i", str(target),
            "-vf", f"fps={fps}",
            "-frames:v", str(max_frames),
            pattern,
        ]

        try:
            proc = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=180,
            )
        except Exception as e:
            return {"ok": False, "data": None, "error": f"extract failed: {e}"}

        if proc.returncode != 0:
            return {
                "ok": False,
                "data": None,
                "error": f"ffmpeg error: {proc.stderr[:400]}",
            }

        frames = sorted(out_dir.glob(f"{target.stem}_*.png"))

        return {
            "ok": True,
            "data": {
                "source": str(target),
                "fps": fps,
                "frame_count": len(frames),
                "frames": [str(p.relative_to(self.workspace_root)) for p in frames],
            },
            "error": None,
        }


TOOLS = [
    ImageDescribe,
    ImageReadText,
    PdfRead,
    DocumentRead,
    VideoRead,
    VideoExtractFrames,
]
